"""Create the six Reading and Writing Skills task instruction and rubric pairs."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'output/reading-writing-assessments'
SOURCE = 'https://lrmds.deped.gov.ph/detail/14520'
TASKS = [
 dict(n=1, title='Building a Coherent School Advisory', codes=['EN11/12RWS-IIIa-1','EN11/12RWS-IIIa-2'],
  goal='Organize related information into a coherent advisory for senior high school learners.',
  prompt='The school reading club needs an advisory inviting learners to a book exchange. The exchange will be held in the library on Friday from 3:00 to 4:00 p.m. Learners may bring one clean, usable book. Participation is voluntary. A learner who cannot bring a book may still join the reading discussion. Books must be checked by the club facilitator before they are exchanged. The advisory must explain the purpose, procedure, and inclusive participation rule.',
  steps=['Create a brainstorming list containing at least eight relevant ideas from the brief. Separate essential details from optional supporting details.',
   'Arrange the essential details in a graphic organizer, then write a topic outline with three main points and at least two supporting points for each.',
   'Write a 250 to 300 word advisory using the brief. Give it a clear purpose, connected paragraphs, and an appropriate closing. Do not add unverified dates, fees, or school policies.',
   'Underline five cohesive devices in your advisory. Add a 100 to 150 word explanation of how the order of ideas and references between sentences make the text connected discourse.'],
  submit='Submit the brainstorming list, graphic organizer, topic outline, advisory, and explanation in one document.',
  criteria=[('Information selection',25,'Includes all event facts and the inclusive participation rule without unsupported additions.'),('Organization',25,'Uses a logical organizer and outline that guide the advisory from purpose to procedure.'),('Coherence and cohesion',30,'Connects ideas clearly and explains five purposeful cohesive devices.'),('Language and mechanics',20,'Uses an appropriate school advisory tone and accurate sentence structure and punctuation.')]),
 dict(n=2,title='Developing an Idea Through Different Writing Patterns',codes=['EN11/12RWS-IIIbf-3'],
  goal='Select and critique patterns of development according to purpose and audience.',
  prompt='Use responsible use of shared school spaces as the common theme. Your audience is fellow senior high school learners. Draw on your observations without naming or identifying individual learners. You may use a clearly identified imagined incident for the narration; do not present invented observations or statistics as research findings.',
  steps=['Prepare eight short paragraphs of 80 to 100 words each: narration, description, definition, exemplification or classification, comparison and contrast, cause and effect, problem and solution, and persuasion. Label each pattern.',
   'For each paragraph, state its purpose in one sentence and underline the details or transitions that establish its pattern. Keep the same overall theme while changing how the information is developed.',
   'Choose two of your paragraphs. Write a 200 to 250 word critique comparing their selection, organization, and development of information. Explain which would work better for a student awareness campaign and why.',
   'Revise the weaker of the two paragraphs using your critique. Submit both its first version and revised version, with a short explanation of the most important change.'],
  submit='Submit the eight paragraphs, the comparison critique, and one revised paragraph with its revision explanation.',
  criteria=[('Accuracy of patterns',35,'Uses all eight patterns distinctly and accurately.'),('Purpose and supporting details',25,'Selects specific details suited to each purpose without unsupported factual claims.'),('Comparative critique',25,'Evaluates organization and development using concrete features from two paragraphs.'),('Revision and mechanics',15,'Makes a purposeful revision and presents readable, edited writing.')]),
 dict(n=3,title='Revising a Student Information Article',codes=['EN11/12RWS-IIIgh-4'],
  goal='Improve organization, coherence, cohesion, language use, and mechanics through revision.',
  prompt='Revise this draft article for learners who are preparing for an independent study session: “Studying is important and students study many things. Put everything on the desk this makes it better. Some people use phones, it is useful and it is distracting. A plan was made by the learner and they do it. First you finish and then you decide what to study. Notes should be checked because those are there. Breaks are nice. There are many lessons and it is difficult so just study them. Students must always know everything after one reading. The place is noisy sometimes. This is why a good plan is needed.”',
  steps=['Annotate at least eight specific problems in the draft. Classify them under organization, coherence and cohesion, language use, or mechanics. Explain why each affects a reader.',
   'Create a sentence outline that orders the information around preparation, focused study, and checking understanding. Replace absolute or unsupported advice with reasonable, qualified statements.',
   'Rewrite the article in 300 to 350 words with a clear title, introduction, connected body paragraphs, and conclusion. Preserve the study planning purpose while developing the missing explanations.',
   'Complete a revision log with at least six before and after entries. Explain the reason for each change. Finish with a 100 word reflection on the difference between revising ideas and proofreading surface errors.'],
  submit='Submit the annotated draft, sentence outline, revised article, revision log, and reflection.',
  criteria=[('Diagnosis of writing problems',25,'Identifies at least eight concrete issues and classifies their effects accurately.'),('Organization and cohesion',30,'Builds a logical article with clear references and transitions.'),('Language and mechanics',25,'Uses precise, qualified language and correct sentence boundaries, spelling, and punctuation.'),('Revision reasoning',20,'Explains substantive changes in the log and distinguishes revision from proofreading.')]),
 dict(n=4,title='Evaluating Claims in a School Reading Proposal',codes=['EN11/12RWS-IIIij-5','EN11/12RWS-IIIij-6','EN11/12RWS-IVac-7'],
  goal='Evaluate explicit and implicit claims and explain how context shapes an argument.',
  prompt='Read the following proposal: “Our class should reserve fifteen minutes each week for voluntary reading. A regular reading period gives learners a predictable opportunity to read beyond their assigned tasks. A quiet classroom is a better place to read than a crowded corridor. Because everyone owns suitable books, the activity will cost nothing. Learners who choose their own reading material are more responsible than learners who do not. The class may borrow books from the library, but the proposal does not explain how borrowing would be arranged. Before making the activity compulsory, representatives should ask learners about access to reading materials and consult the subject teacher about available time.”',
  steps=['Identify six claims in the proposal. For each, state whether it is explicit or implicit, classify it as fact, value, or policy, and justify the classification. A factual claim is not automatically a verified fact.',
   'Choose two claims that need stronger support. Explain what evidence would be needed and whether the passage supplies it. Distinguish textual support from evidence that must be obtained elsewhere.',
   'Locate one relevant school or public education text available to you, such as a library guideline or official reading advisory. Record its title, author or institution, date if available, and URL or file page. Explain one intertextual connection; if using a webpage, explain how a relevant hyperlink affects the reading path.',
   'Write a 600 to 800 word critical response discussing the proposal’s purpose, audience, assumptions, strength of support, and the context supplied by the second text. End with a justified recommendation.'],
  submit='Submit the claim analysis, evidence checks, second source record, and critical response. Include citations for all borrowed wording or ideas.',
  criteria=[('Claim identification',30,'Accurately distinguishes explicit and implicit claims and fact, value, and policy.'),('Evaluation of support',30,'Tests assumptions and identifies specific evidence gaps without inventing data.'),('Context and source connection',20,'Uses a traceable second text and explains intertext or hypertext meaningfully.'),('Reasoned response and writing',20,'Presents a coherent, cited response with a justified recommendation.')]),
 dict(n=5,title='Writing an Evidence Based Critique and Counterclaim',codes=['EN11/12RWS-IVac-8','EN11/12RWS-IVac-9','EN11/12RWS-IVac-10'],
  goal='Develop evaluative statements and counterclaims supported by traceable textual evidence.',
  prompt='Select a nonfiction article, commentary, or school proposal about a learning issue, such as access to reading resources, responsible technology use, or study environments. Use a text that has an identifiable author or institution and can be submitted or linked for verification. Select a second credible text presenting a different perspective. Do not use private messages, classmates’ personal information, or an assessment answer key.',
  steps=['Record the title, author or institution, publication date if available, and URL or file reference for both texts. State each text’s central claim in your own words.',
   'Prepare an evidence map with three evaluative assertions about the main text. For each assertion, identify an exact quotation or accurate paraphrase, its page or paragraph, your interpretation, and a limitation of the evidence.',
   'Formulate two meaningful counterclaims. Respond fairly to each using relevant evidence from the selected texts. Do not distort the opposing position or treat an unsupported opinion as proof.',
   'Write a 1000 word critique, excluding the reference list, evaluating the main text’s claims, context, reasoning, and writing. Integrate your counterclaims and end with a supported overall evaluation. Attach the evidence map and reference list.'],
  submit='Submit the critique, evidence map, reference list, and accessible copies or links to both selected texts.',
  criteria=[('Evaluative reasoning',30,'Makes clear, defensible judgments about claims, context, reasoning, and writing.'),('Evidence and citation',30,'Uses accurate, traceable evidence and explains its relevance and limitations.'),('Counterclaims',25,'Represents two opposing positions fairly and answers them with relevant reasoning.'),('Structure and academic language',15,'Develops a coherent critique with an accurate reference list and edited language.')]),
 dict(n=6,title='Preparing a School Project Proposal and Formal Letter',codes=['EN11/12RWS-IVdg-11','EN11/12RWS-IVdg-12.4','EN11/12RWS-IVhj-13.4'],
  goal='Adapt academic writing and professional correspondence to a shared purpose and audience.',
  prompt='Develop a feasible proposal for a class reading resource exchange or a quiet study corner. Address it to a school coordinator who needs to assess its purpose, beneficiaries, process, resource needs, and safeguards. This is a request for consideration; do not claim that the school has already approved funds, dates, access, or activities.',
  steps=['Create a planning sheet identifying the intended reader, purpose, central need, desired action, and information the reader needs before deciding.',
   'Write a 700 to 900 word project proposal with a rationale, measurable objectives, beneficiaries, activities and timeline, responsibilities, resources and estimated costs, risk controls, and evaluation plan. Explain the basis of any estimates; do not invent research results.',
   'Prepare a one page formal request letter of 200 to 250 words to accompany the proposal. Include a date, recipient designation, subject, salutation, purpose, specific requested action, respectful closing, and your name. Use school contact details only if authorized.',
   'Include an initial outline, a revised version, and a 150 word process note explaining at least three revisions. Check that the letter and proposal agree on the requested action, scope, and resource needs.'],
  submit='Submit the planning sheet, initial outline, final proposal, formal letter, reference list if sources are used, and revision note.',
  criteria=[('Purpose and audience',20,'Adapts the proposal and letter to the coordinator’s decision and uses appropriate conventions.'),('Proposal completeness and feasibility',35,'Provides consistent objectives, process, timeline, responsibilities, resources, safeguards, and evaluation.'),('Professional correspondence',25,'Writes a complete, concise formal letter that requests a specific action and matches the proposal.'),('Writing process and accuracy',20,'Shows purposeful revision, accurate estimates and references, and edited language.')]),
]

def base(title,task):
 d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75)
 for st in d.styles:
  if st.type == 1:
   st.font.color.rgb=RGBColor(0,0,0)
   for color in st.element.xpath('.//w:color'):
    for key in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+key),None)
   for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 s.page_width=Inches(8.27);s.page_height=Inches(11.69)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6)
 d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(13)
 d.core_properties.author='Stanley Vargas Garcia';d.core_properties.title=title
 d.add_paragraph('Interface Computer College Caloocan',style='Subtitle')
 d.add_paragraph(title,style='Title')
 d.add_paragraph(f'Reading and Writing Skills  |  RW-11-12  |  Task {task["n"]:02d}  |  Quarter {3 if task["n"]<=3 else 4}')
 f=s.footer.paragraphs[0];f.add_run('Reading and Writing Skills  |  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 return d

def instructions(t):
 d=base(t['title'],t);d.add_paragraph(t['goal'])
 d.add_paragraph('Name _________________________  Section __________  Date __________')
 d.add_heading('Learning targets',1);d.add_paragraph('Competency alignment: '+', '.join(t['codes'])+'.')
 d.add_heading('Task brief',1);d.add_paragraph(t['prompt'])
 d.add_heading('What to do',1)
 for i,step in enumerate(t['steps'],1):d.add_paragraph(f'{i}. {step}')
 d.add_heading('Submission and scoring',1);d.add_paragraph(t['submit'])
 d.add_paragraph('Submit one DOCX or PDF through this task in SIGMA ELMS. Follow the due date set for your class. Complete your own work, credit all sources, and do not include personal information about others. Total: 100 points. Use the attached rubric to check your work before submission.')
 d.add_paragraph('Curriculum reference: DepEd K to 12 Senior High School Reading and Writing Skills Curriculum Guide, December 2013. '+SOURCE,style='Caption')
 return d

def rubric(t):
 d=base('Scoring Rubric for '+t['title'],t)
 d.add_paragraph('Use this rubric to evaluate the submitted work and give specific feedback. Total available score is 100 points.')
 d.add_paragraph('Learner __________________________  Section __________  Assessor __________________')
 d.add_heading('Criteria and evidence',1)
 for name,weight,evidence in t['criteria']:
  d.add_paragraph(f'{name}  {weight} points',style='Heading 2');d.add_paragraph(evidence)
 d.add_heading('Performance levels',1)
 table=d.add_table(rows=1,cols=3);table.autofit=False
 for c,w in zip(table.columns,[1.1,1.1,4.55]):c.width=Inches(w)
 for cell,txt in zip(table.rows[0].cells,['Level','Multiplier','Evidence across each criterion']):cell.text=txt
 levels=[('4 Accomplished','1.00','Meets all stated evidence for the criterion accurately and consistently; choices are specific and well justified.'),('3 Proficient','0.75','Meets most evidence with minor omissions or errors that do not obscure the main meaning.'),('2 Developing','0.50','Meets some evidence, but important gaps, inconsistencies, or weak explanations limit effectiveness.'),('1 Beginning','0.25','Shows limited relevant evidence with major errors or missing development.'),('0 No evidence','0.00','Provides no assessable evidence for this criterion.')]
 for row in levels:
  for cell,txt in zip(table.add_row().cells,row):cell.text=txt
 for i,row in enumerate(table.rows):
  for cell,width in zip(row.cells,[1.1,1.1,4.55]):
   cell.width=Inches(width)
   tcpr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcpr.append(b)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');tcpr.append(sh)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(4)
    for r in p.runs:r.font.size=Pt(10)
 d.add_paragraph('For each criterion, award its maximum points multiplied by the chosen level multiplier. Add the four criterion scores. Example of calculation only: 25 × 0.75 = 18.75 points. Round the final total to two decimal places.')
 d.add_paragraph('Scores: '+'  |  '.join(f'{x[0]} ____/{x[1]}' for x in t['criteria']))
 d.add_paragraph('Total ____/100     Strength __________________________________________')
 d.add_paragraph('Next step _______________________________________________________')
 return d

OUT.mkdir(parents=True,exist_ok=True)
catalog=json.loads((OUT.parents[1]/'data/curriculum/deped-shs-catalog.json').read_text(encoding='utf-8'))
subject=next(s for s in catalog['subjects'] if s['id']=='CORE-112-RW')
for t,topic in zip(TASKS,subject['topics']):
 t['topicId']=topic['id'];t['topicTitle']=topic['title'];t['quarter']=topic['quarter']
 t['id']=f'RW-11-12-TASK-{t["n"]:02d}'
 for kind,build in [('Instructions',instructions),('Rubric',rubric)]:
  stem=f'RW_Task_{t["n"]:02d}_{kind}';build(t).save(OUT/(stem+'.docx'))
(OUT/'assessment-manifest.json').write_text(json.dumps(TASKS,indent=2),encoding='utf-8')
print(json.dumps({'tasks':len(TASKS),'docxFiles':12,'output':str(OUT)}))
